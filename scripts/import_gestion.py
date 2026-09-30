#!/usr/bin/env python3
"""Construit le registre clients de l'outil à partir du fichier GESTION_COMMERCIALE.

Usage : python3 scripts/import_gestion.py <fichier.xlsx> <dossier_sortie>

Produit dans <dossier_sortie> :
  clients/<CODE>.json  un document par client identifié (collection "clients")
  meta.json            date d'arrêté du fichier, à écrire dans meta/import
  rapport.md           lignes du fichier à renommer, nouveaux clients, cadeaux
Et met à jour scripts/codes_clients.json (codes stables d'une semaine à l'autre).
"""
import datetime
import json
import os
import re
import sys
import unicodedata
import warnings

import openpyxl

ICI = os.path.dirname(os.path.abspath(__file__))
REF = json.load(open(os.path.join(ICI, "referentiel_clients.json"), encoding="utf-8"))
CODES_PATH = os.path.join(ICI, "codes_clients.json")

# Même logique que suggestClientCode() dans l'outil.
TITRES = {"col", "colonel", "mme", "madame", "mr", "monsieur", "m", "dr", "docteur",
          "pr", "professeur", "mlle", "mademoiselle", "cdt", "commandant"}
SAVEURS_2026 = {3: "Passion", 4: "Chocolat", 5: "Citron", 6: "Ananas", 7: "Bissap",
                8: "Crème de Bissap", 9: "Corossol", 10: "Mandarine", 11: "Mangoustan",
                12: "Crème de Cocota"}


def sans_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def cle(nom):
    """Clé de regroupement : même personne si seules la casse, les accents ou les espaces diffèrent."""
    s = sans_accents(str(nom)).upper().strip()
    return re.sub(r"\s+", " ", s)


FUSIONS = {cle(k): v for k, v in REF["fusions"].items()}


def anonyme(k):
    return ("AMBULANT" in k or k.startswith("CLIENT") or k == "XX"
            or re.fullmatch(r"[\d\s.,]+", k) is not None)


def mots_significatifs(nom):
    mots = [w for w in str(nom).split() if w.lower().rstrip(".") not in TITRES]
    mots = [re.sub(r"[^A-Z]", "", sans_accents(w).upper()) for w in mots]
    return [w for w in mots if len(w) >= 3]


def proposer_code(nom, pris):
    m = mots_significatifs(nom) or [re.sub(r"[^A-Z]", "", sans_accents(nom).upper()) or "CLI"]
    candidats = [m[0][:4]]
    if len(m) > 1:
        candidats += [m[0][:3] + m[1][:1], m[0][:2] + m[1][:2], m[1][:4]]
    for c in candidats:
        if c not in pris:
            return c
    base, i = m[0][:3], 2
    while f"{base}{i}" in pris:
        i += 1
    return f"{base}{i}"


def iso(d):
    return d.strftime("%Y-%m-%d") if isinstance(d, datetime.datetime) else None


def main(xlsx, sortie):
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    groupes = {}   # clé canonique -> infos

    def groupe(nom_brut):
        k = cle(nom_brut)
        if k in FUSIONS:
            nom = FUSIONS[k]
            kc = cle(nom)
        else:
            if anonyme(k):
                return None
            nom, kc = None, k
        g = groupes.setdefault(kc, {"impose": nom, "orth": {}, "bouteilles": 0, "achats": 0,
                                    "premiere": None, "derniere": None, "cadeaux": [],
                                    "lignes_variantes": []})
        return g

    def noter(g, nom_brut, d, feuille, ligne):
        brut = str(nom_brut).strip()
        g["orth"][brut] = g["orth"].get(brut, 0) + (2 if feuille == "VENTES 2026" else 1)
        g["premiere"] = min(filter(None, [g["premiere"], d]))
        g["derniere"] = max(filter(None, [g["derniere"], d]))
        g["lignes_variantes"].append((feuille, ligne, brut))

    cadeaux_manuels = {(cle(c["client"]), c["date"]): c["libelle"] for c in REF["cadeaux"]}

    ws = wb["VENTES 2026"]
    arrete = None
    for r in range(4, ws.max_row + 1):
        d, nom = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not isinstance(d, datetime.datetime) or not nom:
            continue
        arrete = max(filter(None, [arrete, d]))
        g = groupe(nom)
        if g is None:
            continue
        noter(g, nom, d, "VENTES 2026", r)
        qte = ws.cell(r, 13).value or 0
        net = ws.cell(r, 18).value
        kc = cle(g["impose"] or nom)
        manuel = cadeaux_manuels.get((kc, iso(d)))
        if manuel or net == 0:
            produits = [SAVEURS_2026[c] for c in SAVEURS_2026 if ws.cell(r, c).value]
            g["cadeaux"].append({"date": iso(d),
                                 "libelle": manuel or ("Bouteille offerte — " + ", ".join(produits))})
            continue
        g["bouteilles"] += qte
        g["achats"] += 1

    ws = wb["VENTES 2025"]
    for r in range(4, ws.max_row + 1):
        d, nom = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(d, datetime.datetime) and nom:
            g = groupe(nom)
            if g is not None:
                noter(g, nom, d, "VENTES 2025", r)

    # Ancienneté déclarée dans HISTORIQUE 10 ANS : prime sur les feuilles de vente.
    depuis_hist = {}
    ws = wb["HISTORIQUE 10 ANS"]
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip().upper() == "CLIENT" and \
                str(ws.cell(r, 2).value or "").strip().upper() == "DEPUIS":
            for rr in range(r + 1, r + 40):
                nom, an = ws.cell(rr, 1).value, ws.cell(rr, 2).value
                if nom and an and re.fullmatch(r"\d{4}", str(an).strip()):
                    g = groupe(nom)
                    if g is not None:
                        depuis_hist[id(g)] = str(an).strip()
                        g["orth"].setdefault(str(nom).strip(), 0)

    codes = json.load(open(CODES_PATH, encoding="utf-8")) if os.path.exists(CODES_PATH) else {}
    for nom, code in REF["codes_imposes"].items():
        codes.setdefault(nom, code)
    pris = set(codes.values())

    os.makedirs(os.path.join(sortie, "clients"), exist_ok=True)
    rapport_renommer, nouveaux, docs = [], [], []
    for kc, g in sorted(groupes.items()):
        if g["impose"]:
            nom = g["impose"]
        else:
            nom = max(g["orth"].items(), key=lambda x: (x[1], x[0]))[0]
        code = codes.get(nom) or codes.get(cle(nom))
        if not code:
            code = proposer_code(nom, pris)
            codes[nom] = code
            nouveaux.append(nom)
        pris.add(code)
        for feuille, ligne, brut in g["lignes_variantes"]:
            if feuille == "VENTES 2026" and cle(brut) != cle(nom):
                rapport_renommer.append((feuille, ligne, brut, nom))
        depuis = depuis_hist.get(id(g)) or iso(g["premiere"]) or ""
        doc = {
            "nom": nom, "code": code, "cle": cle(nom),
            "variantes": sorted(({cle(o) for o in g["orth"]}
                                 | {k for k, v in FUSIONS.items() if v == nom}) - {cle(nom)}),
            "annee": arrete.year if arrete else datetime.date.today().year,
            "bouteilles": int(g["bouteilles"]), "achats": g["achats"],
            "depuis": depuis, "derniereCommande": iso(g["derniere"]) or "",
            "cadeauxFichier": g["cadeaux"],
        }
        docs.append(doc)
        with open(os.path.join(sortie, "clients", code + ".json"), "w", encoding="utf-8") as f:
            json.dump(doc, f, ensure_ascii=False, indent=1)

    json.dump(dict(sorted(codes.items())), open(CODES_PATH, "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    meta = {"fichier": os.path.basename(xlsx), "arreteAu": iso(arrete),
            "importeLe": datetime.date.today().isoformat(), "nbClients": len(docs)}
    json.dump(meta, open(os.path.join(sortie, "meta.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    with open(os.path.join(sortie, "rapport.md"), "w", encoding="utf-8") as f:
        f.write(f"# Import {meta['fichier']} — arrêté au {meta['arreteAu']}\n\n")
        f.write(f"{len(docs)} clients identifiés, {len(nouveaux)} nouveaux codes.\n\n")
        f.write("## Lignes à renommer dans VENTES 2026\n\n")
        for feuille, ligne, brut, nom in rapport_renommer:
            f.write(f"- ligne {ligne} : « {brut} » → « {nom} »\n")
        f.write("\n## Cadeaux (exclus du cumul)\n\n")
        for d in docs:
            for c in d["cadeauxFichier"]:
                f.write(f"- {d['nom']} : {c['libelle']} ({c['date']})\n")
        f.write("\n## Nouveaux clients\n\n" + "".join(f"- {n} → {codes[n]}\n" for n in nouveaux))
    print(open(os.path.join(sortie, "rapport.md"), encoding="utf-8").read()[:3000])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
